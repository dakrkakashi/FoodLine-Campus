import sys
import pptx
sys.stdout.reconfigure(encoding='utf-8')

prs = pptx.Presentation(r'C:\Users\shiva\Desktop\main finance\ppt.pptx')
for idx, slide in enumerate(prs.slides, 1):
    for s in slide.shapes:
        if s.has_table:
            t = s.table
            print(f'\nSlide {idx} Table ({len(t.rows)} rows x {len(t.columns)} cols):')
            for r_idx in range(len(t.rows)):
                cell0 = t.cell(r_idx, 0)
                fill_col = cell0.fill.fore_color.rgb if cell0.fill and cell0.fill.type == 1 else 'None'
                txt = cell0.text.strip()
                p0 = cell0.text_frame.paragraphs[0] if cell0.text_frame.paragraphs else None
                font_bold = p0.font.bold if p0 and p0.font else None
                print(f'  Row {r_idx:02d}: fill={fill_col}, bold={font_bold}, text="{txt[:35]}"')
